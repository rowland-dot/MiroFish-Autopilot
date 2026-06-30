"""Render report Markdown into a downloadable .docx (python-docx).

Powers the report's Word download. Handles the constructs the report agent
emits: ATX headings, paragraphs, **bold** spans, and `-`/`*` bullet lists.
Sets a CJK-friendly font so Chinese renders cleanly in Word.
"""

import io
import re

from docx import Document
from docx.oxml.ns import qn

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*]\s+(.*)$")
_RULE = re.compile(r"^---+$")
_BOLD = re.compile(r"\*\*(.+?)\*\*")

_CJK_FONT = "Microsoft YaHei"


def _apply_cjk_font(doc: Document) -> None:
    """Make default + heading styles render CJK text cleanly in Word."""
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        try:
            style = doc.styles[style_name]
            style.font.name = _CJK_FONT
            rpr = style.element.get_or_add_rPr()
            rfonts = rpr.get_or_add_rFonts()
            rfonts.set(qn("w:eastAsia"), _CJK_FONT)
        except KeyError:
            continue


def _add_inline(paragraph, text: str) -> None:
    """Add text to a paragraph, rendering **bold** spans as bold runs."""
    pos = 0
    for match in _BOLD.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        paragraph.add_run(match.group(1)).bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def markdown_to_docx_bytes(md: str) -> bytes:
    """Convert report Markdown into .docx file bytes."""
    doc = Document()
    _apply_cjk_font(doc)

    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip() or _RULE.match(line):
            continue

        heading = _HEADING.match(line)
        if heading:
            level = min(len(heading.group(1)), 4)
            doc.add_heading(heading.group(2).strip(), level=level)
            continue

        bullet = _BULLET.match(line)
        if bullet:
            para = doc.add_paragraph(style="List Bullet")
            _add_inline(para, bullet.group(1).strip())
            continue

        _add_inline(doc.add_paragraph(), line.strip())

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
