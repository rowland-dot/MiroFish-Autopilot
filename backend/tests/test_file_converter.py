"""TDD: uploaded document files are converted to Markdown before ingest.

The converter is the single entry point for turning a non-text source file
(.docx, and other office/pdf formats markitdown handles) into Markdown that
the existing GraphRAG text pipeline consumes.
"""
from docx import Document

from app.utils.file_converter import convert_to_markdown


def test_convert_docx_preserves_heading_and_body_as_markdown(tmp_path):
    # Arrange: a .docx with a Heading 1 and two body paragraphs (one Chinese)
    docx_path = tmp_path / "product.docx"
    doc = Document()
    doc.add_heading("Product Overview", level=1)
    doc.add_paragraph("Tropical Oasis brain-power supplement.")
    doc.add_paragraph("控糖人群友好，零糖配方。")
    doc.save(str(docx_path))

    # Act
    md = convert_to_markdown(str(docx_path))

    # Assert: real Markdown output, not plain text — heading becomes an ATX
    # heading and all body text (incl. Chinese) survives.
    assert "# Product Overview" in md
    assert "Tropical Oasis brain-power supplement." in md
    assert "控糖人群友好，零糖配方。" in md
