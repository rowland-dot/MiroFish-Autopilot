"""TDD: FileParser routes .docx through the markdown converter.

The upload pipeline calls FileParser.extract_text() on each saved file. For
.docx it must return Markdown (via file_converter), so the rest of the ingest
flow is unchanged. .pdf / .md / .txt behaviour is untouched.
"""
from docx import Document

from app.utils.file_parser import FileParser


def test_docx_is_supported():
    assert FileParser.is_supported("brief.docx") is True


def test_extract_text_converts_docx_to_markdown(tmp_path):
    p = tmp_path / "brief.docx"
    d = Document()
    d.add_heading("Brand Brief", level=1)
    d.add_paragraph("Premium positioning, 控糖 friendly.")
    d.save(str(p))

    text = FileParser.extract_text(str(p))

    assert "# Brand Brief" in text
    assert "Premium positioning, 控糖 friendly." in text


def test_markdown_extraction_unchanged(tmp_path):
    # Regression: existing .md path must keep returning raw text verbatim.
    # write_bytes avoids Windows text-mode newline translation in the fixture.
    p = tmp_path / "notes.md"
    p.write_bytes(b"# Title\n\nplain body")
    assert FileParser.extract_text(str(p)) == "# Title\n\nplain body"
