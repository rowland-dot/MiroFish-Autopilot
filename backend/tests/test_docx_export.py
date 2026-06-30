"""TDD: render report Markdown into a downloadable .docx (python-docx).

Drives the Word download: given the report's stored markdown_content, produce
.docx bytes whose headings, body text (incl. Chinese), bold, and bullets are
preserved with real Word styles.
"""
import io

from docx import Document

from app.utils.docx_export import markdown_to_docx_bytes


def _paras(data: bytes):
    doc = Document(io.BytesIO(data))
    return [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]


def test_headings_body_chinese_and_bullets_are_rendered():
    md = (
        "# Report Title\n\n"
        "Intro paragraph.\n\n"
        "## Section One\n\n"
        "控糖人群友好，零糖配方。\n\n"
        "- bullet one\n"
        "- bullet two\n"
    )
    paras = _paras(markdown_to_docx_bytes(md))

    assert ("Heading 1", "Report Title") in paras
    assert ("Heading 2", "Section One") in paras
    assert any(t == "Intro paragraph." for _, t in paras)
    assert any("控糖人群友好" in t for _, t in paras)
    assert any(s == "List Bullet" and t == "bullet one" for s, t in paras)


def test_inline_bold_becomes_a_bold_run():
    data = markdown_to_docx_bytes("**Key finding:** strong demand.")
    doc = Document(io.BytesIO(data))
    para = next(p for p in doc.paragraphs if p.text.strip())
    bold_runs = [r.text for r in para.runs if r.bold]
    assert "Key finding:" in bold_runs


def test_output_is_a_valid_docx_zip():
    # .docx is a zip; first bytes are the PK zip signature.
    data = markdown_to_docx_bytes("# Hi")
    assert data[:2] == b"PK"
